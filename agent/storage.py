import os
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from agent.models import FlightOffer

COLUMNS = [
    "recorded_at", "airline", "price", "currency",
    "departure", "arrival", "duration", "stops", "origin", "destination",
]
COLUMNS_RT = [
    "recorded_at", "airline", "price", "currency",
    "departure", "arrival", "return_date", "duration", "stops", "origin", "destination",
]


def _sheet_name(origin: str, destination: str, roundtrip: bool = False) -> str:
    sep = "⇄" if roundtrip else "→"
    return f"{origin.upper()}{sep}{destination.upper()}"[:31]  # Excel sheet name limit


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


class PriceStorage:
    def __init__(self, data_dir: str = "data"):
        self._path = Path(data_dir) / "flight_prices.xlsx"
        self._path.parent.mkdir(parents=True, exist_ok=True)

    def _load_or_create_workbook(self) -> Workbook:
        if self._path.exists():
            return openpyxl.load_workbook(self._path)
        wb = Workbook()
        wb.remove(wb.active)  # remove default blank sheet
        return wb

    def _get_or_create_sheet(self, wb: Workbook, name: str, columns: list[str]):
        if name in wb.sheetnames:
            return wb[name]
        ws = wb.create_sheet(name)
        ws.append(columns)
        _style_header(ws)
        for i, _ in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
        return ws

    def append_offers(self, offers: list[FlightOffer]) -> None:
        if not offers:
            return
        wb = self._load_or_create_workbook()
        # Group by route + trip type (round-trip offers carry a return_date)
        groups: dict[tuple, list[FlightOffer]] = {}
        for o in offers:
            groups.setdefault((o.origin, o.destination, o.return_date is not None), []).append(o)

        for (origin, destination, roundtrip), items in groups.items():
            columns = COLUMNS_RT if roundtrip else COLUMNS
            ws = self._get_or_create_sheet(wb, _sheet_name(origin, destination, roundtrip), columns)
            for o in items:
                head = [
                    o.recorded_at.strftime("%Y-%m-%d %H:%M:%S"),
                    o.airline, o.price, o.currency,
                    o.departure_time.strftime("%Y-%m-%d %H:%M"),
                    o.arrival_time.strftime("%Y-%m-%d %H:%M"),
                ]
                tail = [o.duration, o.stops, o.origin, o.destination]
                ws.append(head + ([o.return_date] + tail if roundtrip else tail))
        wb.save(self._path)

    def load_history(self, origin: str, destination: str, roundtrip: bool = False) -> list[dict]:
        if not self._path.exists():
            return []
        wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        try:
            name = _sheet_name(origin, destination, roundtrip)
            if name not in wb.sheetnames:
                return []
            rows = list(wb[name].iter_rows(values_only=True))
        finally:
            wb.close()  # read_only mode keeps the file open until closed (locks it on Windows)
        if len(rows) < 2:
            return []
        headers = rows[0]
        return [dict(zip(headers, row)) for row in rows[1:]]

    @property
    def path(self) -> Path:
        return self._path
